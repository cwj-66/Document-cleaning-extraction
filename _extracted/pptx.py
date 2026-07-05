import re
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from itertools import groupby
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER


def clean_text(text: str) -> str:
    '''Tab 转空格，合并连续空白，去首尾空白。'''
    text = text.replace('\t', ' ')
    text = re.sub(r'\s+', ' ', text)
    return text.strip()


# PPT 的 Element 比 DOCX 多一个 slide_num 字段（幻灯片页码）
@dataclass
class Element:
    '''解析出的一个幻灯片元素。'''
    index:      int
    category:   str   # Title / Text / Table / Image
    text:       str   # 清洗后的文字；Image 格式为 "[图片: fname | 页: N: 标题]"
    slide_num:  int = 0    # 所在幻灯片页码（从 1 起）—— PPT 特有
    image_path: str = ''   # 仅 Image 类型有值

# DOCX 靠 Header/Footer/PageBreak 类型过滤元信息
# PPT 没有这些类型，改用「频次检测」：在 ≥60% 的页面出现 → 视为模板噪音

# 标题占位符类型（居中标题 or 普通标题）
_TITLE_PH_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}


def _get_slide_title(slide) -> str:
    '''从标题占位符读取该页标题文字。'''
    for shape in slide.shapes:
        if shape.is_placeholder and shape.placeholder_format.type in _TITLE_PH_TYPES:
            t = clean_text(shape.text_frame.text)
            if t:
                return t
    return ''


def _detect_boilerplate(prs: Presentation, threshold: float = 0.6) -> set[str]:
    '''
    检测模板噪音文本：在超过 threshold 比例的幻灯片上出现的短文本片段。
    常见：公司名、页脚日期、版权声明、每页都有的 slogan。
    '''
    text_slide_count: Counter = Counter()
    total = len(prs.slides)
    for slide in prs.slides:
        seen: set[str] = set()
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                t = clean_text(para.text)
                # 只考虑短文本（长段落不可能是模板噪音）
                if t and len(t) <= 60:
                    seen.add(t)
        text_slide_count.update(seen)
    return {t for t, n in text_slide_count.items() if n / total >= threshold}


def _table_to_text(table) -> str:
    '''把 pptx Table 对象转为纯文本（每行用 " | " 分隔列）。'''
    rows = []
    for row in table.rows:
        cells = [clean_text(cell.text_frame.text) for cell in row.cells]
        rows.append(' | '.join(cells))
    return '\n'.join(rows)


def extract_slides(pptx_path: Path) -> tuple[list[Element], set[str]]:
    '''
    解析 pptx，返回 (Element 列表, 检测到的模板噪音集合)。

    每个 Element 都带 slide_num（页码）。
    图片文件保存到 {stem}_images/ 目录，text 格式：[图片: fname | 页: N: 标题]。
    '''
    prs = Presentation(str(pptx_path))
    boilerplate = _detect_boilerplate(prs)

    img_dir = pptx_path.parent / f'{pptx_path.stem}_images'
    img_dir.mkdir(exist_ok=True)

    elements: list[Element] = []

    for slide_idx, slide in enumerate(prs.slides, start=1):
        slide_title = _get_slide_title(slide)
        page_label  = f'{slide_idx}: {slide_title}' if slide_title else str(slide_idx)
        img_counter = 0

        for shape in slide.shapes:

            # 图片 —— PPT 中图片是独立 Picture 形状，不像 DOCX 嵌在段落里
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                img_counter += 1
                img_ext   = shape.image.ext or 'png'
                img_fname = f'slide{slide_idx:02d}_img{img_counter}.{img_ext}'
                (img_dir / img_fname).write_bytes(shape.image.blob)
                elements.append(Element(
                    index=len(elements),
                    category='Image',
                    text=f'[图片: {img_fname} | 页: {page_label}]',
                    slide_num=slide_idx,
                    image_path=str(img_dir / img_fname),
                ))
                continue

            # 表格
            if shape.has_table:
                t = _table_to_text(shape.table)
                if t.strip():
                    elements.append(Element(
                        index=len(elements),
                        category='Table',
                        text=t,
                        slide_num=slide_idx,
                    ))
                continue

            # 文本框
            if not shape.has_text_frame:
                continue

            is_title = (
                shape.is_placeholder
                and shape.placeholder_format.type in _TITLE_PH_TYPES
            )

            for para in shape.text_frame.paragraphs:
                t = clean_text(para.text)
                if not t or t in boilerplate:
                    continue
                elements.append(Element(
                    index=len(elements),
                    category='Title' if is_title else 'Text',
                    text=t,
                    slide_num=slide_idx,
                ))

    return elements, boilerplate

@dataclass
class RagChunk:
    '''RAG 用的一块数据，带元信息方便检索。'''
    text:        str
    source_file: str
    section:     str
    chunk_index: int
    slide_num:   int = 0   # PPT 特有
    char_count:  int = field(init=False)

    def __post_init__(self):
        self.char_count = len(self.text)


import base64
import os

from openai import OpenAI

_CAPTION_MODEL = 'qwen3.6-flash'


def caption_image(img_path: str) -> str:
    '''调百炼多模态模型生成图片描述。'''
    client = OpenAI(
        api_key=os.environ['DASHSCOPE_API_KEY'],
        base_url='https://dashscope.aliyuncs.com/compatible-mode/v1',
    )
    ext = Path(img_path).suffix.lstrip('.') or 'jpeg'
    b64 = base64.b64encode(Path(img_path).read_bytes()).decode()
    resp = client.chat.completions.create(
        model=_CAPTION_MODEL,
        messages=[{
            'role': 'user',
            'content': [
                {'type': 'image_url', 'image_url': {'url': f'data:image/{ext};base64,{b64}'}},
                {'type': 'text', 'text': '请用中文描述这张图片的内容，包括文字和结构。'},
            ],
        }],
    )
    return resp.choices[0].message.content


def chunk_slides(
    elements:   list[Element],
    source_file: str,
    chunk_size:  int = 600,
    min_chunk_size: int = 10,
) -> list[RagChunk]:
    '''
    PPT 分块：默认 1 页 = 1 块（标题 + 全部正文文本合并）。

    - Title  → section 标题，加 "# " 前缀加入文本
    - Text   → 追加到当前页文本
    - Table  → flush 当前页文本，表格单独成块
    - Image  → flush 当前页文本，图片单独成块
    - 合并文本超过 chunk_size → 按字数平切（无 overlap）
    '''
    chunks: list[RagChunk] = []

    def add_chunk(text: str, section: str, slide_num: int):
        if len(text.strip()) >= min_chunk_size:
            chunks.append(RagChunk(
                text=text.strip(),
                source_file=source_file,
                section=section,
                chunk_index=len(chunks),
                slide_num=slide_num,
            ))

    def flush_text(text_parts: list[str], section: str, slide_num: int):
        if not text_parts:
            return
        combined = '\n'.join(text_parts)
        if len(combined) <= chunk_size:
            add_chunk(combined, section, slide_num)
        else:
            # 超长则平切，PPT 单页很少出现此情况
            for i in range(0, len(combined), chunk_size):
                add_chunk(combined[i:i + chunk_size], section, slide_num)

    # 按 slide_num 分组（提取时已按页顺序，sort 保险）
    sorted_els = sorted(elements, key=lambda e: e.slide_num)

    for slide_num, group in groupby(sorted_els, key=lambda e: e.slide_num):
        slide_els = list(group)

        # 确定该页 section（取标题占位符文字）
        title_el = next((e for e in slide_els if e.category == 'Title'), None)
        section  = title_el.text if title_el else f'第{slide_num}页'

        text_parts: list[str] = []

        for el in slide_els:
            if el.category == 'Image':
                flush_text(text_parts, section, slide_num)
                text_parts = []
                text = el.text
                if el.image_path:
                    text += f'\n描述: {caption_image(el.image_path)}'
                add_chunk(text, section, slide_num)
            elif el.category == 'Table':
                flush_text(text_parts, section, slide_num)
                text_parts = []
                add_chunk(el.text, section, slide_num)
            elif el.category == 'Title':
                text_parts.append(f'# {el.text}')
            else:
                text_parts.append(el.text)

        flush_text(text_parts, section, slide_num)

    return chunks

rag_chunks = chunk_slides(elements, source_file=pptx_path.name, chunk_size=600)

text_chunks  = [c for c in rag_chunks if not c.text.startswith('[图片:')]
image_chunks = [c for c in rag_chunks if c.text.startswith('[图片:')]

print(f'总元素: {len(elements)} 个 → RAG 分块: {len(rag_chunks)} 块')
print(f'  文字块: {len(text_chunks)}  图片块: {len(image_chunks)}')
print(f'总字符: {sum(c.char_count for c in rag_chunks)}')
print('=' * 60)
for c in rag_chunks:
    preview = c.text[:70] + '...' if len(c.text) > 70 else c.text
    tag = '[图]' if c.text.startswith('[图片:') else '   '
    print(f'\n{tag} [{c.chunk_index:2d}] p{c.slide_num:02d} section={c.section[:25]:25s} len={c.char_count:4d}')
    print(f'    {preview}')