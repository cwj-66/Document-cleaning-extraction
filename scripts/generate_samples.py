"""Generate fictional local examples, without downloading third-party content.

Basic mode needs only the standard library. --documents additionally creates
DOCX, PPTX, XLSX, image and PDF examples using the optional document libraries.
Existing files are left untouched unless --overwrite is explicitly supplied.
"""

import argparse
import csv
from email.message import EmailMessage
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
ROWS = [["Item", "Quantity", "Note"], ["Sample A", "12", "Fictional data"], ["Sample B", "8", "For parser demonstration"]]
TEXT = (
    "This document contains fictional information for a document cleaning example. "
    "Sample A has twelve units and Sample B has eight units. "
    "Keep the source file and section when splitting content into chunks. "
)


def target(relative: str, overwrite: bool) -> Path | None:
    path = ROOT / relative
    if path.exists() and not overwrite:
        return None
    path.parent.mkdir(parents=True, exist_ok=True)
    return path


def basic(overwrite: bool) -> None:
    path = target("workshops/html/.html/sample.html", overwrite)
    if path:
        path.write_text(
            '<!doctype html><html><head><meta charset="utf-8"><title>文档清洗示例</title></head>'
            '<body><nav>导航：这部分应被过滤</nav><h1>示例知识文档</h1>'
            '<p>这是一份虚构资料，用于演示正文提取和按章节分块，不含真实业务信息。</p>'
            '<h2>物品说明</h2><ul><li>示例 A：数量 12</li><li>示例 B：数量 8</li></ul>'
            '<table><tr><th>物品</th><th>数量</th></tr><tr><td>示例 A</td><td>12</td></tr></table>'
            '<footer>页脚：这部分应被过滤</footer></body></html>\n', encoding="utf-8",
        )
    path = target("workshops/email/sample.eml", overwrite)
    if path:
        message = EmailMessage()
        message["Subject"] = "文档清洗示例"
        message["From"] = "sender@example.com"
        message["To"] = "reader@example.com"
        message["Date"] = "Sun, 04 Oct 2026 09:00:00 +0800"
        message.set_content(
            "一、当前进度\n这封邮件只有虚构的演示内容，用于检查正文提取与分块。\n"
            "二、下一步\n- 核对来源信息\n- 检查分块结果\n"
            "\n--\nExample signature\n\nOn Sat, someone wrote:\n> Older reply for removal.\n"
        )
        path.write_bytes(message.as_bytes())
    for extension, separator in [("csv", ","), ("tsv", "\t")]:
        path = target(f"workshops/xlsx/sample.{extension}", overwrite)
        if path:
            with path.open("w", encoding="utf-8-sig", newline="") as stream:
                csv.writer(stream, delimiter=separator).writerows(ROWS)


def documents(overwrite: bool) -> None:
    import fitz
    from docx import Document
    from openpyxl import Workbook
    from PIL import Image, ImageDraw
    from pptx import Presentation

    path = target("workshops/word/.docx/sample.docx", overwrite)
    if path:
        doc = Document()
        doc.add_heading("Document cleaning sample", 0)
        doc.add_heading("1. Overview", 1)
        doc.add_paragraph(TEXT * 3)
        doc.add_heading("2. Inventory", 1)
        table = doc.add_table(rows=0, cols=3)
        for row in ROWS:
            cells = table.add_row().cells
            for cell, value in zip(cells, row):
                cell.text = value
        doc.save(path)
    path = target("workshops/ppt/.ppt/sample.pptx", overwrite)
    if path:
        deck = Presentation()
        for title, content in [("Document cleaning sample", TEXT), ("Keep useful context", TEXT * 2)]:
            slide = deck.slides.add_slide(deck.slide_layouts[1])
            slide.shapes.title.text = title
            slide.placeholders[1].text = content
        deck.save(path)
    path = target("workshops/xlsx/sample.xlsx", overwrite)
    if path:
        book = Workbook()
        book.active.title = "Inventory"
        for row in ROWS:
            book.active.append(row)
        book.create_sheet("Empty")
        book.save(path)
    # Render our own PDF as a scan so OCR has genuine raster input.
    with fitz.open() as document:
        for index in range(2):
            page = document.new_page()
            page.insert_text((50, 60), f"Sample section {index + 1}", fontsize=20)
            page.insert_textbox(fitz.Rect(50, 95, 530, 600), TEXT * 5, fontsize=12)
        path = target("workshops/pdf/纯文本.pdf/sample-text.pdf", overwrite)
        if path:
            document.save(path)
        with fitz.open() as scan:
            for page in document:
                image = page.get_pixmap(matrix=fitz.Matrix(2, 2))
                scan.new_page(width=page.rect.width, height=page.rect.height).insert_image(page.rect, stream=image.tobytes("png"))
            path = target("workshops/pdf/扫描.pdf/sample-scan.pdf", overwrite)
            if path:
                scan.save(path)
    image = Image.new("RGB", (1200, 700), "white")
    draw = ImageDraw.Draw(image)
    # Large text makes this generated image useful for OCR smoke checks.
    from PIL import ImageFont
    font = ImageFont.load_default(size=30)
    for index, row in enumerate(ROWS):
        draw.text((50, 50 + index * 100), "    |    ".join(row), fill="black", font=font)
    path = target("workshops/img/sample.png", overwrite)
    if path:
        image.save(path)
    with fitz.open() as mixed:
        page = mixed.new_page()
        page.insert_text((50, 60), "Illustrated inventory sample", fontsize=20)
        page.insert_textbox(fitz.Rect(50, 95, 530, 270), TEXT * 2, fontsize=12)
        import io
        stream = io.BytesIO()
        image.save(stream, format="PNG")
        page.insert_image(fitz.Rect(50, 300, 530, 580), stream=stream.getvalue())
        path = target("workshops/pdf/图文混排.pdf/sample-mixed.pdf", overwrite)
        if path:
            mixed.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--documents", action="store_true", help="also generate document, image and PDF samples")
    parser.add_argument("--overwrite", action="store_true", help="replace existing generated samples")
    args = parser.parse_args()
    basic(args.overwrite)
    if args.documents:
        documents(args.overwrite)
    print("Generated fictional samples. Audio input must be provided separately.")


if __name__ == "__main__":
    main()
